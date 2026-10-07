"""Conversion functions. Each takes input paths + an output dir and returns the output file path."""
import html
import io
import os
import shutil
import tempfile
import zipfile

import markdown
import mammoth
import openpyxl
import pymupdf
from PIL import Image, ImageOps
from pdf2docx import Converter
from pillow_heif import register_heif_opener

register_heif_opener()  # lets Pillow open .heic/.heif (iPhone photos)

MARGINS = {"none": 0, "small": 28, "normal": 54, "large": 90}  # points (72 = 1in)


def page_layout(page_size="a4", orientation="portrait", margin="normal", **_):
    rect = pymupdf.paper_rect(page_size if page_size in ("a3", "a4", "a5", "letter", "legal") else "a4")
    if orientation == "landscape":
        rect = pymupdf.Rect(0, 0, rect.height, rect.width)
    return rect, MARGINS.get(margin, 54)
BASE_CSS = """
body { font-family: sans-serif; font-size: 11pt; line-height: 1.4; }
pre { font-family: monospace; font-size: 9.5pt; white-space: pre-wrap; }
table { border-collapse: collapse; }
td, th { border: 1px solid #999; padding: 3px 6px; font-size: 9pt; }
th { background-color: #eee; }
img { max-width: 100%; }
"""

IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".jfif", ".webp", ".bmp", ".gif", ".tif", ".tiff",
              ".heic", ".heif", ".avif", ".ico"}
DOCUMENT_EXTS = {".svg", ".epub", ".xps", ".oxps", ".fb2", ".mobi", ".cbz"}


def _stem(path):
    return os.path.splitext(os.path.basename(path))[0]


def _zip(paths, out_path):
    with zipfile.ZipFile(out_path, "w", zipfile.ZIP_DEFLATED) as z:
        for p in paths:
            z.write(p, os.path.basename(p))
    return out_path


# ---------------------------------------------------------------- PDF -> X

def pdf_to_docx(src, out_dir, **_):
    out = os.path.join(out_dir, _stem(src) + ".docx")
    cv = Converter(src)
    try:
        cv.convert(out)
    finally:
        cv.close()
    return out


PIL_FORMATS = {"webp": "WEBP", "tiff": "TIFF", "bmp": "BMP", "gif": "GIF", "avif": "AVIF"}


def pdf_to_images(src, out_dir, fmt="png", dpi=150, **_):
    pages = []
    with pymupdf.open(src) as doc:
        for i, page in enumerate(doc, 1):
            pix = page.get_pixmap(dpi=int(dpi))
            p = os.path.join(out_dir, f"{_stem(src)}_page{i:03d}.{fmt}")
            if fmt in ("jpg", "jpeg"):
                pix.save(p, output="jpg", jpg_quality=90)
            elif fmt in PIL_FORMATS:
                pix.pil_save(p, format=PIL_FORMATS[fmt])
            else:
                pix.save(p)
            pages.append(p)
    if len(pages) == 1:
        return pages[0]
    return _zip(pages, os.path.join(out_dir, f"{_stem(src)}_{fmt}.zip"))


def pdf_to_svg(src, out_dir, **_):
    pages = []
    with pymupdf.open(src) as doc:
        for i, page in enumerate(doc, 1):
            p = os.path.join(out_dir, f"{_stem(src)}_page{i:03d}.svg")
            with open(p, "w", encoding="utf-8") as f:
                f.write(page.get_svg_image())
            pages.append(p)
    if len(pages) == 1:
        return pages[0]
    return _zip(pages, os.path.join(out_dir, f"{_stem(src)}_svg.zip"))


def pdf_to_pptx(src, out_dir, dpi=150, notes=True, **_):
    """One slide per page, rendered as an image so it looks identical; page text goes in the speaker notes."""
    from pptx import Presentation
    from pptx.util import Emu

    emu = lambda pt: Emu(int(pt * 12700))  # 1 pt = 12700 EMU
    out = os.path.join(out_dir, _stem(src) + ".pptx")
    prs = Presentation()
    with pymupdf.open(src) as doc:
        first = doc[0].rect
        prs.slide_width, prs.slide_height = emu(first.width), emu(first.height)
        for page in doc:
            slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank layout
            pix = page.get_pixmap(dpi=int(dpi))
            # pages of a different shape are fitted and centred
            r = min(first.width / page.rect.width, first.height / page.rect.height)
            w, h = page.rect.width * r, page.rect.height * r
            slide.shapes.add_picture(io.BytesIO(pix.tobytes("png")), emu((first.width - w) / 2),
                                     emu((first.height - h) / 2), emu(w), emu(h))
            text = page.get_text().strip()
            if notes and text:
                slide.notes_slide.notes_text_frame.text = text
    prs.save(out)
    return out


def pdf_to_txt(src, out_dir, **_):
    out = os.path.join(out_dir, _stem(src) + ".txt")
    with pymupdf.open(src) as doc, open(out, "w", encoding="utf-8") as f:
        for i, page in enumerate(doc, 1):
            if i > 1:
                f.write("\n\f\n")
            f.write(page.get_text())
    return out


def pdf_to_html(src, out_dir, **_):
    out = os.path.join(out_dir, _stem(src) + ".html")
    with pymupdf.open(src) as doc:
        body = "\n".join(page.get_text("xhtml") for page in doc)
    with open(out, "w", encoding="utf-8") as f:
        f.write(f"<!doctype html><html><head><meta charset='utf-8'>"
                f"<title>{html.escape(_stem(src))}</title></head><body>\n{body}\n</body></html>")
    return out


def pdf_to_xlsx(src, out_dir, **_):
    """Detected tables become sheets; if none are found, each page's text lines go in one sheet."""
    out = os.path.join(out_dir, _stem(src) + ".xlsx")
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    with pymupdf.open(src) as doc:
        for pno, page in enumerate(doc, 1):
            for tno, table in enumerate(page.find_tables().tables, 1):
                ws = wb.create_sheet(f"P{pno} T{tno}")
                for row in table.extract():
                    ws.append([c if c is not None else "" for c in row])
        if not wb.sheetnames:
            ws = wb.create_sheet("Text")
            for pno, page in enumerate(doc, 1):
                for line in page.get_text().splitlines():
                    ws.append([pno, line])
    wb.save(out)
    return out


# ---------------------------------------------------------------- X -> PDF

def html_to_pdf_file(html_str, out, archive=None, **layout):
    page, margin = page_layout(**layout)
    story = pymupdf.Story(html=html_str, user_css=BASE_CSS, archive=archive)
    writer = pymupdf.DocumentWriter(out)
    where = page + (margin, margin, -margin, -margin)
    more = True
    while more:
        dev = writer.begin_page(page)
        more, _ = story.place(where)
        story.draw(dev)
        writer.end_page()
    writer.close()
    return out


def document_to_pdf(src, out_dir, **_):
    """SVG, EPUB, XPS, FB2, MOBI, CBZ — formats MuPDF opens natively."""
    out = os.path.join(out_dir, _stem(src) + ".pdf")
    with pymupdf.open(src) as doc, pymupdf.open("pdf", doc.convert_to_pdf()) as pdf:
        pdf.save(out, garbage=3, deflate=True)
    return out


def images_to_pdf(srcs, out_dir, page_size="fit", orientation="auto", margin="none", **_):
    """page_size 'fit' makes each page exactly the image size; otherwise images are centred on paper."""
    out = os.path.join(out_dir, (_stem(srcs[0]) if len(srcs) == 1 else "images") + ".pdf")
    doc = pymupdf.open()
    for s in srcs:
        img = Image.open(s)
        try:
            for i in range(getattr(img, "n_frames", 1)):  # multi-page TIFF / animated GIF
                img.seek(i)
                frame = ImageOps.exif_transpose(img)
                alpha = frame.mode in ("RGBA", "LA") or "transparency" in frame.info
                buf = io.BytesIO()
                if alpha:
                    frame.convert("RGBA").save(buf, "PNG")
                else:
                    frame.convert("RGB").save(buf, "JPEG", quality=92)
                w, h = frame.size
                if page_size == "fit":
                    pw, ph, m = w * 72 / 150, h * 72 / 150, MARGINS.get(margin, 0)  # 150 dpi
                    rect = pymupdf.Rect(0, 0, pw + 2 * m, ph + 2 * m)
                else:
                    o = orientation if orientation != "auto" else ("landscape" if w > h else "portrait")
                    rect, m = page_layout(page_size, o, margin)
                page = doc.new_page(width=rect.width, height=rect.height)
                page.insert_image(page.rect + (m, m, -m, -m), stream=buf.getvalue(), keep_proportion=True)
        finally:
            img.close()
    doc.save(out, garbage=3, deflate=True)
    doc.close()
    return out


def txt_to_pdf(src, out_dir, **layout):
    with open(src, encoding="utf-8", errors="replace") as f:
        text = f.read()
    return html_to_pdf_file(f"<pre>{html.escape(text)}</pre>",
                            os.path.join(out_dir, _stem(src) + ".pdf"), **layout)


def md_to_pdf(src, out_dir, **layout):
    with open(src, encoding="utf-8", errors="replace") as f:
        body = markdown.markdown(f.read(), extensions=["tables", "fenced_code"])
    return html_to_pdf_file(body, os.path.join(out_dir, _stem(src) + ".pdf"),
                            archive=os.path.dirname(src), **layout)


def html_to_pdf(src, out_dir, **layout):
    with open(src, encoding="utf-8", errors="replace") as f:
        body = f.read()
    return html_to_pdf_file(body, os.path.join(out_dir, _stem(src) + ".pdf"),
                            archive=os.path.dirname(src), **layout)


def _docx_to_pdf_fallback(src, out, **layout):
    img_dir = tempfile.mkdtemp()
    counter = [0]

    def save_image(image):
        counter[0] += 1
        name = f"img{counter[0]}.{image.content_type.split('/')[-1]}"
        with image.open() as data, open(os.path.join(img_dir, name), "wb") as f:
            shutil.copyfileobj(data, f)
        return {"src": name}

    try:
        with open(src, "rb") as f:
            body = mammoth.convert_to_html(f, convert_image=mammoth.images.img_element(save_image)).value
        return html_to_pdf_file(body, out, archive=img_dir, **layout)
    finally:
        shutil.rmtree(img_dir, ignore_errors=True)


def _xlsx_to_pdf_fallback(src, out, **layout):
    wb = openpyxl.load_workbook(src, data_only=True, read_only=True)
    parts = []
    for ws in wb.worksheets:
        rows = "".join(
            "<tr>" + "".join(f"<td>{html.escape('' if v is None else str(v))}</td>" for v in row) + "</tr>"
            for row in ws.iter_rows(values_only=True)
        )
        parts.append(f"<h2>{html.escape(ws.title)}</h2><table>{rows}</table>")
    wb.close()
    return html_to_pdf_file("".join(parts), out, **layout)


def office_to_pdf(src, out_dir, **layout):
    """Word / Excel / PowerPoint -> PDF. Uses installed MS Office for full fidelity, else a Python fallback."""
    from office import export_pdf, office_available

    ext = os.path.splitext(src)[1].lower()
    out = os.path.join(out_dir, _stem(src) + ".pdf")
    if office_available():
        try:
            return export_pdf(src, out)
        except Exception:
            if ext in (".ppt", ".pptx", ".doc", ".xls"):
                raise  # no fallback for these
    if ext == ".docx":
        return _docx_to_pdf_fallback(src, out, **layout)
    if ext == ".xlsx":
        return _xlsx_to_pdf_fallback(src, out, **layout)
    raise RuntimeError(f"Converting {ext} to PDF requires Microsoft Office.")
